#!/usr/bin/env python3
"""
수치해석 기말 프로젝트 발표 PPT 생성
python3 make_ppt.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import os

# ── 색상 팔레트 ────────────────────────────────────────────────────────────────
C_DARK   = RGBColor(0x1A, 0x1A, 0x2E)   # 진남색 (배경)
C_BLUE   = RGBColor(0x16, 0x21, 0x3E)   # 슬라이드 배경
C_ACCENT = RGBColor(0x0F, 0x3C, 0x96)   # 강조 파랑
C_GREEN  = RGBColor(0x00, 0xB4, 0x50)   # 인라이어 초록
C_RED    = RGBColor(0xE7, 0x2B, 0x2B)   # 아웃라이어 빨강
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_LGRAY  = RGBColor(0xD0, 0xD8, 0xE8)
C_YELLOW = RGBColor(0xFF, 0xD7, 0x00)
C_ORANGE = RGBColor(0xFF, 0x6B, 0x35)

W, H = Inches(13.33), Inches(7.5)   # 와이드 16:9

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]   # 완전 빈 레이아웃


# ── 공통 헬퍼 ──────────────────────────────────────────────────────────────────
def bg(slide, color=C_BLUE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, l, t, w, h, text='', font_size=18, bold=False,
        fg=C_WHITE, bg_color=None, align=PP_ALIGN.LEFT,
        border_color=None, border_width=Pt(0)):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    if bg_color:
        txBox.fill.solid()
        txBox.fill.fore_color.rgb = bg_color
    if border_color:
        txBox.line.color.rgb = border_color
        txBox.line.width = border_width
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(font_size)
    run.font.bold  = bold
    run.font.color.rgb = fg
    return txBox

def hline(slide, t, color=C_ACCENT, width_in=12.0, l=0.67):
    shape = slide.shapes.add_shape(
        1, Inches(l), Inches(t), Inches(width_in), Inches(0.04))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def para(slide, l, t, w, h, lines, sizes, bolds=None, colors=None,
         align=PP_ALIGN.LEFT, spacing=1.15):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    bolds  = bolds  or [False] * len(lines)
    colors = colors or [C_WHITE] * len(lines)
    for i, (text, size, bold, color) in enumerate(zip(lines, sizes, bolds, colors)):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * (spacing - 1) * 0.7)
        run = p.add_run()
        run.text = text
        run.font.size  = Pt(size)
        run.font.bold  = bold
        run.font.color.rgb = color

def rect(slide, l, t, w, h, fill=C_ACCENT, alpha=None):
    shape = slide.shapes.add_shape(
        1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape

def img(slide, path, l, t, w, h=None):
    if not os.path.exists(path):
        return
    if h:
        slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))
    else:
        slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w))


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 1: 표지
# ══════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
bg(s1, C_DARK)
rect(s1, 0, 0, 13.33, 0.08, C_ACCENT)
rect(s1, 0, 7.42, 13.33, 0.08, C_ACCENT)

# 장식 사각형
rect(s1, 0.5, 1.5, 0.12, 4.0, C_ACCENT)

para(s1, 0.85, 1.4, 11.5, 1.5,
     ['수치해석 기말 프로젝트'],
     [22], [False], [C_LGRAY])

para(s1, 0.85, 2.05, 11.5, 2.2,
     ['3D LiDAR 포인트 클라우드에서\nRANSAC + 최소제곱법을 이용한\n바닥 평면 탐지'],
     [36], [True], [C_WHITE])

hline(s1, 4.55, C_ACCENT, 11.0, 0.85)

para(s1, 0.85, 4.75, 11.5, 2.0,
     ['청주대학교  ·  수치해석  ·  2025년'],
     [16], [False], [C_LGRAY])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 2: 목차
# ══════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
bg(s2)
rect(s2, 0, 0, 13.33, 1.1, C_ACCENT)
para(s2, 0.5, 0.22, 12.0, 0.8, ['목  차'], [32], [True], [C_WHITE])

items = [
    ('01', '연구 배경 및 목적',         '왜 바닥 탐지가 어려운가?'),
    ('02', '수업 내용 → 3D 확장',      '1D 직선 피팅 → 3D 평면 피팅'),
    ('03', '알고리즘 구성',              'RANSAC + 최소제곱법 (정규방정식)'),
    ('04', '실험 환경',                  'Gazebo 시뮬레이션 + Livox Mid-360'),
    ('05', '시연 및 결과',               'RViz 시각화 · 지표 비교'),
    ('06', '결론',                       ''),
]

for i, (num, title, sub) in enumerate(items):
    t = 1.35 + i * 0.97
    rect(s2, 0.5, t, 0.55, 0.72, C_ACCENT)
    para(s2, 0.53, t + 0.12, 0.5, 0.5, [num], [18], [True], [C_WHITE], PP_ALIGN.CENTER)
    para(s2, 1.22, t + 0.05, 10.5, 0.75,
         [title, sub] if sub else [title],
         [18, 13] if sub else [18],
         [True, False] if sub else [True],
         [C_WHITE, C_LGRAY] if sub else [C_WHITE])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 3: 연구 배경 및 목적
# ══════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
bg(s3)
rect(s3, 0, 0, 13.33, 1.1, C_ACCENT)
para(s3, 0.5, 0.22, 12.0, 0.8, ['01  연구 배경 및 목적'], [28], [True], [C_WHITE])

# 왼쪽 설명
para(s3, 0.5, 1.3, 6.5, 5.5,
     ['문제 상황',
      '',
      'LiDAR(레이저 거리 센서)가 공간을 스캔하면',
      '바닥, 벽, 상자 등이 모두 섞인',
      '"포인트 클라우드"가 생성됩니다.',
      '',
      '이 중에서 바닥 포인트만 골라',
      '평면을 정확히 탐지하는 것이 목표입니다.',
      '',
      '평가 기준',
      '',
      '• 각도 오차 (Angle Error)',
      '  탐지 평면이 수평(0°)에서 얼마나 틀어졌는가',
      '',
      '• RMSE',
      '  바닥 포인트들이 탐지 평면에서 평균적으로',
      '  얼마나 떨어져 있는가',
      '',
      '• 인라이어 비율 (IR)',
      '  전체 포인트 중 바닥으로 분류된 비율',
      ],
     [16, 6, 14, 14, 14, 6, 14, 14, 6,
      16, 6, 14, 13, 6, 14, 13, 13, 6, 14, 13],
     [True, False, False, False, False, False, False, False, False,
      True, False, False, False, False, False, False, False, False, False, False],
     [C_YELLOW, C_WHITE, C_WHITE, C_WHITE, C_WHITE, C_WHITE, C_WHITE, C_WHITE, C_WHITE,
      C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY, C_WHITE, C_LGRAY, C_LGRAY, C_LGRAY, C_WHITE, C_LGRAY, C_LGRAY])

# 오른쪽: 어려운 이유 박스
rect(s3, 7.3, 1.3, 5.6, 5.5, RGBColor(0x0A, 0x18, 0x38))
para(s3, 7.5, 1.5, 5.2, 5.0,
     ['왜 어려운가?',
      '',
      '❌  단순 LS 적용 시',
      '    벽/박스 포인트가 평면을 위로 끌어올림',
      '    → 바닥 평면이 기울어짐',
      '',
      '✅  RANSAC + LS 적용 시',
      '    아웃라이어 제거 후 바닥만 선별',
      '    → 정확한 수평 평면 탐지',
      '',
      '실험 환경 (Gazebo)',
      '',
      '  바닥 50×50m  +  외벽 4개',
      '  L자 벽  +  박스 2개  +  기둥',
      '  경사판(15°)',
      '',
      '  → 전체 포인트의 ~50%가 아웃라이어',
      ],
     [16, 6, 14, 13, 13, 6, 14, 13, 13, 6,
      16, 6, 13, 13, 13, 6, 14],
     [True, False, True, False, False, False, True, False, False, False,
      True, False, False, False, False, False, False],
     [C_YELLOW, C_WHITE, C_WHITE, C_LGRAY, C_LGRAY, C_WHITE,
      C_WHITE, C_LGRAY, C_LGRAY, C_WHITE,
      C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY, C_LGRAY, C_WHITE, C_WHITE])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 4: 수업 내용 → 3D 확장
# ══════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
bg(s4)
rect(s4, 0, 0, 13.33, 1.1, C_ACCENT)
para(s4, 0.5, 0.22, 12.0, 0.8, ['02  수업 내용 → 3D 확장'], [28], [True], [C_WHITE])

# 왼쪽: 1D (수업)
rect(s4, 0.4, 1.25, 5.8, 5.6, RGBColor(0x0A, 0x18, 0x38))
para(s4, 0.6, 1.4, 5.4, 5.2,
     ['수업에서 배운 것  (1D 직선 피팅)',
      '',
      '모델:   y = a₀ + a₁x',
      '',
      '목적:   Sr = Σ(yᵢ - a₀ - a₁xᵢ)²  최소화',
      '',
      '풀이:   ∂Sr/∂a₀ = 0',
      '        ∂Sr/∂a₁ = 0',
      '',
      '정규방정식 (2×2):',
      '',
      '  ⎡  n    Σxᵢ  ⎤ ⎡a₀⎤   ⎡ Σyᵢ  ⎤',
      '  ⎢             ⎥ ⎢  ⎥ = ⎢      ⎥',
      '  ⎣ Σxᵢ  Σxᵢ² ⎦ ⎣a₁⎦   ⎣ Σxᵢyᵢ⎦',
      '',
      '미지수 2개 (a₀, a₁)',
      ],
     [15, 5, 14, 5, 14, 5, 14, 14, 5, 14, 5,
      12, 12, 12, 5, 13],
     [True, False, False, False, False, False, False, False, False,
      True, False, False, False, False, False, False],
     [C_YELLOW] + [C_WHITE]*15)

# 화살표
para(s4, 6.35, 3.6, 0.7, 0.8, ['→'], [32], [True], [C_ACCENT], PP_ALIGN.CENTER)

# 오른쪽: 3D (프로젝트)
rect(s4, 7.1, 1.25, 5.8, 5.6, RGBColor(0x0A, 0x18, 0x38))
para(s4, 7.3, 1.4, 5.4, 5.2,
     ['프로젝트에서 사용  (3D 평면 피팅)',
      '',
      '모델:   z = a₀ + a₁x + a₂y',
      '',
      '목적:   Sr = Σ(zᵢ - a₀ - a₁xᵢ - a₂yᵢ)²  최소화',
      '',
      '풀이:   ∂Sr/∂a₀ = 0',
      '        ∂Sr/∂a₁ = 0      (변수 하나 추가)',
      '        ∂Sr/∂a₂ = 0',
      '',
      '정규방정식 (3×3):',
      '',
      '  ⎡  n    Σxᵢ   Σyᵢ  ⎤ ⎡a₀⎤   ⎡  Σzᵢ  ⎤',
      '  ⎢ Σxᵢ  Σxᵢ²  Σxᵢyᵢ⎥ ⎢a₁⎥ = ⎢ Σxᵢzᵢ ⎥',
      '  ⎣ Σyᵢ  Σxᵢyᵢ Σyᵢ² ⎦ ⎣a₂⎦   ⎣ Σyᵢzᵢ ⎦',
      '',
      '미지수 3개 (a₀, a₁, a₂)',
      ],
     [15, 5, 14, 5, 14, 5, 14, 14, 14, 5, 14, 5,
      12, 12, 12, 5, 13],
     [True, False, False, False, False, False, False, False, False,
      False, True, False, False, False, False, False, False],
     [C_ORANGE] + [C_WHITE]*16)

# 하단 요약
rect(s4, 0.4, 6.95, 12.5, 0.42, C_ACCENT)
para(s4, 0.6, 7.0, 12.2, 0.38,
     ['핵심: 변수가 x 하나 → x, y 둘로 늘었을 뿐, 구조(Sr 편미분 → 정규방정식 → 행렬 풀기)는 동일'],
     [13], [True], [C_WHITE])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 5: 알고리즘 구성
# ══════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
bg(s5)
rect(s5, 0, 0, 13.33, 1.1, C_ACCENT)
para(s5, 0.5, 0.22, 12.0, 0.8, ['03  알고리즘 구성'], [28], [True], [C_WHITE])

# RANSAC 박스
rect(s5, 0.4, 1.25, 5.9, 5.1, RGBColor(0x0A, 0x18, 0x38))
rect(s5, 0.4, 1.25, 5.9, 0.45, RGBColor(0x0F, 0x3C, 0x96))
para(s5, 0.5, 1.3, 5.7, 0.4, ['① Custom RANSAC'], [15], [True], [C_WHITE])

para(s5, 0.55, 1.82, 5.6, 4.4,
     ['목적: 아웃라이어(벽/박스)를 제거하고',
      '      바닥 포인트만 골라내기',
      '',
      '알고리즘 (500회 반복):',
      '',
      '  STEP 1.  3개 포인트 랜덤 선택',
      '',
      '  STEP 2.  외적(cross product)으로',
      '           법선벡터 → 평면 방정식 결정',
      '           n = (p2-p1) × (p3-p1)',
      '',
      '  STEP 3.  나머지 전체 포인트에서',
      '           평면까지 거리 < 5cm인 것',
      '           → 인라이어로 분류',
      '',
      '  STEP 4.  인라이어 수 최대인 평면 보존',
      '',
      '최종 출력: 최적 평면 + 인라이어 집합',
      ],
     [13]*18,
     [False, False, False, True, False,
      True, False, True, False, False, False,
      True, False, False, False,
      True, False, True],
     [C_WHITE]*3 + [C_YELLOW] + [C_WHITE] +
     [C_LGRAY, C_WHITE, C_LGRAY, C_WHITE, C_WHITE, C_WHITE,
      C_LGRAY, C_WHITE, C_WHITE, C_WHITE,
      C_LGRAY, C_WHITE, C_GREEN])

# 화살표
para(s5, 6.4, 3.5, 0.6, 0.8, ['+'], [28], [True], [C_ACCENT], PP_ALIGN.CENTER)

# LS 박스
rect(s5, 7.1, 1.25, 5.8, 5.1, RGBColor(0x0A, 0x18, 0x38))
rect(s5, 7.1, 1.25, 5.8, 0.45, RGBColor(0x7A, 0x2D, 0x00))
para(s5, 7.2, 1.3, 5.6, 0.4, ['② RANSAC + Least Squares (LS)'], [15], [True], [C_WHITE])

para(s5, 7.25, 1.82, 5.5, 4.4,
     ['목적: RANSAC 인라이어 전체를 이용해',
      '      평면을 더 정밀하게 다듬기',
      '',
      '알고리즘:',
      '',
      '  INPUT: RANSAC 인라이어 포인트들',
      '',
      '  설계행렬 A 구성:',
      '  ⎡1  x₁  y₁⎤',
      '  ⎢1  x₂  y₂⎥  (n×3 행렬)',
      '  ⎣...       ⎦',
      '',
      '  정규방정식 풀기:',
      '  (AᵀA) c = Aᵀz',
      '  → np.linalg.solve(AᵀA, Aᵀz)',
      '',
      '  OUTPUT: [a₀, a₁, a₂]',
      '  → z = a₀ + a₁x + a₂y',
      ],
     [13]*18,
     [False, False, False, True, False,
      True, False, True, False, False, False, False,
      True, False, False, False,
      True, False],
     [C_WHITE]*3 + [C_YELLOW] + [C_WHITE] +
     [C_LGRAY, C_WHITE, C_LGRAY, C_WHITE, C_WHITE, C_WHITE, C_WHITE,
      C_LGRAY, C_WHITE, C_WHITE, C_WHITE,
      C_ORANGE, C_WHITE])

# 하단 요약
rect(s5, 0.4, 6.45, 12.5, 0.88, RGBColor(0x0A, 0x18, 0x38))
para(s5, 0.6, 6.5, 12.1, 0.8,
     ['RANSAC → 아웃라이어 제거 (강건성)     +     LS → 전체 인라이어 활용 (정밀도)',
      'IR: 동일 수준  /  RMSE: 40% 감소  /  Angle Error: 75% 감소'],
     [14, 13], [True, False], [C_WHITE, C_LGRAY])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 6: 실험 환경
# ══════════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(BLANK)
bg(s6)
rect(s6, 0, 0, 13.33, 1.1, C_ACCENT)
para(s6, 0.5, 0.22, 12.0, 0.8, ['04  실험 환경'], [28], [True], [C_WHITE])

# 왼쪽 설명
para(s6, 0.5, 1.3, 5.5, 5.8,
     ['시뮬레이션 환경',
      '',
      '  플랫폼     Gazebo Classic 11',
      '  ROS 버전  ROS2 Humble',
      '  LiDAR     Livox Mid-360 (시뮬)',
      '  토픽       /mid360_PointCloud2',
      '',
      '월드 구성  (50m × 50m)',
      '',
      '  ✅  수평 바닥 (탐지 대상, z = 0)',
      '  ❌  외벽 4개 (높이 4m)',
      '  ❌  L자 벽 구조물',
      '  ❌  박스 2개 + 원기둥',
      '  ❌  경사판 (pitch 15°)',
      '',
      '알고리즘 파라미터',
      '',
      '  RANSAC 반복 횟수   500회',
      '  인라이어 임계값    5 cm',
      '  처리 포인트 수     최대 15,000개',
      ],
     [15, 5, 13, 13, 13, 13, 5,
      15, 5, 13, 13, 13, 13, 13, 5,
      15, 5, 13, 13, 13],
     [True] + [False]*19,
     [C_YELLOW] + [C_WHITE]*6 + [C_YELLOW] + [C_WHITE]*6 + [C_WHITE] + [C_YELLOW] + [C_WHITE]*3)

# 오른쪽 통계 박스
rect(s6, 6.3, 1.3, 6.6, 5.8, RGBColor(0x0A, 0x18, 0x38))
para(s6, 6.5, 1.5, 6.2, 5.4,
     ['실제 측정 데이터 (1 프레임)',
      '',
      '  총 포인트 수         9,438 개',
      '  바닥 인라이어        4,743 개  (50.3%)',
      '  아웃라이어           4,695 개  (49.7%)',
      '',
      '30프레임 평균 측정 결과',
      '',
      '                Custom     RANSAC',
      '               RANSAC        + LS',
      '',
      '  IR          0.505±0.006   0.503±0.006',
      '  RMSE    1.09±0.25 cm   0.65±0.05 cm',
      '  Angle   0.095±0.016°   0.024±0.009°',
      '',
      '  RMSE 개선:   40%',
      '  Angle 개선:  75%',
      ],
     [15, 5, 13, 13, 13, 5, 15, 5, 13, 13, 5,
      13, 13, 13, 5, 13, 13],
     [True] + [False]*16,
     [C_YELLOW] + [C_WHITE]*5 + [C_YELLOW] + [C_WHITE]*6 +
     [C_WHITE, C_WHITE, C_WHITE, C_WHITE, C_WHITE] +
     [C_GREEN, C_GREEN])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 7: 시연 및 결과 (그림)
# ══════════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(BLANK)
bg(s7)
rect(s7, 0, 0, 13.33, 1.1, C_ACCENT)
para(s7, 0.5, 0.22, 12.0, 0.8, ['05  시연 및 결과'], [28], [True], [C_WHITE])

fig_path = '/home/seo/ros2_ws/ugv_ws/fig_ls_comparison.png'
if os.path.exists(fig_path):
    img(s7, fig_path, 0.3, 1.2, 12.7, 5.6)
else:
    para(s7, 0.5, 3.0, 12.0, 1.0,
         ['[fig_ls_comparison.png 파일을 여기에 삽입하세요]'],
         [16], [False], [C_LGRAY])

para(s7, 0.5, 6.9, 12.3, 0.5,
     ['초록 = 바닥 인라이어 (4,743pts)     빨강 = 아웃라이어 (4,695pts)     '
      '파란선 = RANSAC 평면     주황선 = RANSAC+LS 평면'],
     [12], [False], [C_LGRAY])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 8: 결과 비교표
# ══════════════════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(BLANK)
bg(s8)
rect(s8, 0, 0, 13.33, 1.1, C_ACCENT)
para(s8, 0.5, 0.22, 12.0, 0.8, ['05  결과 비교'], [28], [True], [C_WHITE])

# 표 헤더
rect(s8, 1.0, 1.3, 11.0, 0.6, C_ACCENT)
headers = ['지표', 'Custom RANSAC', 'RANSAC + LS', '개선율']
lefts   = [1.05, 4.3, 7.8, 11.2]
for h, l in zip(headers, lefts):
    para(s8, l, 1.35, 2.8, 0.5, [h], [14], [True], [C_WHITE], PP_ALIGN.CENTER)

rows = [
    ('Inlier Ratio (IR)', '0.505 ± 0.006', '0.503 ± 0.006', '—'),
    ('RMSE',              '1.09 ± 0.25 cm', '0.65 ± 0.05 cm', '↓ 40%'),
    ('Angle Error',       '0.095 ± 0.016°', '0.024 ± 0.009°', '↓ 75%'),
]
row_colors = [
    [C_WHITE,  C_WHITE,  C_WHITE,  C_LGRAY],
    [C_WHITE,  C_WHITE,  C_GREEN,  C_GREEN],
    [C_WHITE,  C_WHITE,  C_GREEN,  C_GREEN],
]

for ri, (row, rcolors) in enumerate(zip(rows, row_colors)):
    t = 1.98 + ri * 0.75
    bg_c = RGBColor(0x0A, 0x18, 0x38) if ri % 2 == 0 else RGBColor(0x0D, 0x20, 0x45)
    rect(s8, 1.0, t, 11.0, 0.72, bg_c)
    for val, l, rc in zip(row, lefts, rcolors):
        para(s8, l, t + 0.18, 2.8, 0.45, [val], [14], [False], [rc], PP_ALIGN.CENTER)

# 하단 해석
rect(s8, 1.0, 4.28, 11.0, 2.6, RGBColor(0x0A, 0x18, 0x38))
para(s8, 1.2, 4.45, 10.6, 2.3,
     ['해석',
      '',
      '• IR이 두 방법에서 동일  →  RANSAC이 아웃라이어 제거 성능을 결정',
      '   LS는 이미 골라진 인라이어를 더 정밀하게 맞추는 역할',
      '',
      '• RMSE 40% 감소 + 표준편차도 0.25 → 0.05 cm로 감소',
      '   LS가 인라이어 전체 평균을 사용하므로 프레임마다 안정적',
      '',
      '• Angle Error 75% 감소  →  수업 정규방정식의 3D 확장이 실제로 유효함을 확인',
      ],
     [15, 5, 13, 13, 5, 13, 13, 5, 13],
     [True, False, False, False, False, False, False, False, False],
     [C_YELLOW, C_WHITE, C_WHITE, C_LGRAY, C_WHITE, C_WHITE, C_LGRAY, C_WHITE, C_WHITE])


# ══════════════════════════════════════════════════════════════════════════════
# 슬라이드 9: 결론
# ══════════════════════════════════════════════════════════════════════════════
s9 = prs.slides.add_slide(BLANK)
bg(s9, C_DARK)
rect(s9, 0, 0, 13.33, 1.1, C_ACCENT)
para(s9, 0.5, 0.22, 12.0, 0.8, ['06  결론'], [28], [True], [C_WHITE])
rect(s9, 0.5, 1.25, 0.12, 5.5, C_ACCENT)

para(s9, 0.85, 1.3, 11.5, 5.5,
     ['수업 내용의 직접 확장',
      '',
      '    수업에서 배운 1D 정규방정식을 변수 하나만 추가해 3D 평면 피팅에 그대로 적용했습니다.',
      '    코드 6줄이 수업 슬라이드의 행렬식과 1:1 대응됩니다.',
      '',
      '아웃라이어 50% 환경에서도 탐지 성공',
      '',
      '    전체 포인트의 절반이 벽/장애물인 어려운 환경에서',
      '    RANSAC + LS 조합으로 각도 오차 0.024°를 달성했습니다.',
      '',
      'LS 추가로 정밀도 향상',
      '',
      '    RANSAC 단독 대비 RMSE 40%, 각도 오차 75% 감소.',
      '    표준편차도 함께 감소 → 프레임마다 일관된 결과.',
      '',
      '실제 응용',
      '',
      '    드론 착륙 지점 탐지, 자율주행 차량 지면 인식 등',
      '    아웃라이어가 많은 실제 환경에서도 적용 가능한 방법입니다.',
      ],
     [17, 5, 13, 13, 5,
      17, 5, 13, 13, 5,
      17, 5, 13, 13, 5,
      17, 5, 13, 13],
     [True, False, False, False, False,
      True, False, False, False, False,
      True, False, False, False, False,
      True, False, False, False],
     [C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY, C_WHITE,
      C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY, C_WHITE,
      C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY, C_WHITE,
      C_YELLOW, C_WHITE, C_LGRAY, C_LGRAY])

rect(s9, 0, 7.1, 13.33, 0.4, C_ACCENT)
para(s9, 0.5, 7.13, 12.5, 0.35,
     ['감사합니다'],
     [16], [True], [C_WHITE], PP_ALIGN.CENTER)


# ── 저장 ──────────────────────────────────────────────────────────────────────
out = '/home/seo/ros2_ws/ugv_ws/presentation.pptx'
prs.save(out)
print(f'저장 완료: {out}')
print(f'슬라이드 수: {len(prs.slides)}')
